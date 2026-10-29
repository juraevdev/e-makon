"""E-MAKON to'liq loyiha hujjatini Word (.docx) formatida yig'adi.

Ishga tushirish:  python docs/build_loyha_hujjati.py
Natija:           docs/E-MAKON_Loyiha_va_sahifalar.docx

Rasmlar: pitch/shots/ (mavjud bo'lsa qo'shiladi).
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent / "E-MAKON_Loyiha_va_sahifalar.docx"
# Agar asosiy fayl Word/IDE da ochiq bo'lsa, zaxira nomga yoziladi.
OUT_FALLBACK = Path(__file__).resolve().parent / "E-MAKON_Loyiha_toliq.docx"
SHOTS = ROOT / "pitch" / "shots"
ASSETS = ROOT / "pitch" / "assets"

GREEN = RGBColor(0x1E, 0x6B, 0x34)
DARK = RGBColor(0x1F, 0x23, 0x21)
GREY = RGBColor(0x5F, 0x66, 0x62)
HEADER_FILL = "1E6B34"
ZEBRA_FILL = "EEF5EF"

doc = Document()


def setup_styles() -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color in (
        ("Heading 1", 16, GREEN),
        ("Heading 2", 13, GREEN),
        ("Heading 3", 11.5, DARK),
    ):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(6)
    for sec in doc.sections:
        sec.page_height = Cm(29.7)
        sec.page_width = Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(1.8)
        sec.bottom_margin = Cm(1.8)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number_footer() -> None:
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("E-MAKON · Loyiha hujjati · sahifa ")
        run.font.size = Pt(8.5)
        run.font.color.rgb = GREY
        r = p.add_run()
        r.font.size = Pt(8.5)
        r.font.color.rgb = GREY
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            r._r.append(el)


def h1(text: str) -> None:
    doc.add_heading(text, level=1)


def h2(text: str) -> None:
    doc.add_heading(text, level=2)


def h3(text: str) -> None:
    doc.add_heading(text, level=3)


def para(text: str, *, bold_prefix: str | None = None, italic: bool = False) -> None:
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.color.rgb = DARK
    r = p.add_run(text)
    r.italic = italic
    r.font.color.rgb = DARK


def bullets(items) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        if isinstance(item, tuple):
            r = p.add_run(item[0])
            r.bold = True
            p.add_run(item[1])
        else:
            p.add_run(str(item))


def table(headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, HEADER_FILL)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = t.rows[r_i + 1].cells[c_i]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
            if r_i % 2 == 1:
                shade(cell, ZEBRA_FILL)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


def add_image(path: Path, *, width_in: float = 5.8, caption: str = "") -> None:
    if not path.exists():
        para(f"[Rasm topilmadi: {path.name}]", italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width_in))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GREY


def add_image_pair(
    left: Path,
    right: Path,
    *,
    width_in: float = 3.0,
    caption: str = "",
) -> None:
    """Ikki rasmni yonma-yon (jadval ichida)."""
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, path in enumerate((left, right)):
        cell = t.rows[0].cells[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if path.exists():
            run = cell.paragraphs[0].add_run()
            run.add_picture(str(path), width=Inches(width_in))
        else:
            cell.paragraphs[0].add_run(f"[{path.name}]")
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
    doc.add_paragraph()


def cover() -> None:
    for _ in range(2):
        doc.add_paragraph()
    cover_img = ASSETS / "cover.jpg"
    if cover_img.exists():
        add_image(cover_img, width_in=5.5, caption="")
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("E-MAKON")
    r.bold = True
    r.font.size = Pt(36)
    r.font.color.rgb = GREEN

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("Bog‘dorchilik va makon parvarishi platformasi")
    r.font.size = Pt(16)
    r.font.color.rgb = DARK

    d = doc.add_paragraph()
    d.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = d.add_run(
        "Loyiha tavsifi · Nima uchun kerak · Texnologiyalar ·\n"
        "Sahifalar (mobil, firma, superadmin) · Skrinshotlar · Oldindagi ishlar"
    )
    r.font.size = Pt(12)
    r.font.color.rgb = GREY

    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = m.add_run(f"\n{date.today().strftime('%d.%m.%Y')}")
    r.font.size = Pt(11)
    r.font.color.rgb = GREY

    doc.add_page_break()


def build() -> None:
    setup_styles()
    add_page_number_footer()
    cover()

    # ------------------------------------------------------------------ 1
    h1("1. Loyiha haqida")
    para(
        "E-MAKON — uy-joy atrofi, bog‘, landshaft va yillik parvarish xizmatlarini "
        "mijozlar, firmalar va platforma ma’muriyatini bir joyga jamlovchi raqamli bozor (marketplace)."
    )
    para(
        "Mijoz mobil ilova yoki Telegram orqali buyurtma beradi. Firma o‘z panelidan "
        "buyurtmani qabul qiladi, xodim va jadvalni boshqaradi. Superadmin esa sifat, "
        "to‘lov (escrow), moderatsiya va butun platformani nazorat qiladi."
    )
    bullets(
        [
            ("Shior: ", "«Makoningiz — bir platformada»."),
            ("Maqsad: ", "Ishonchsiz telefon kelishuvlari o‘rniga shaffof, kuzatiladigan va himoyalangan xizmat."),
            ("Kanallar: ", "Mobil ilova (Flutter), Telegram bot, firma admin paneli, superadmin paneli, yagona backend API."),
            ("Interfeys tili: ", "Foydalanuvchi interfeysi o‘zbek tilida (lotin)."),
        ]
    )

    # ------------------------------------------------------------------ 2
    h1("2. Nima uchun kerak? (ahamiyat)")
    problem = ASSETS / "problem.jpg"
    if problem.exists():
        add_image(problem, width_in=5.5, caption="1-rasm. Muammo: tartibsiz xizmat bozori va ishonchsizlik")

    h2("2.1. Bugungi muammo")
    bullets(
        [
            "Mijoz ishonchli usta/firma topa olmaydi — telefon orqali tasodifiy kelishuv.",
            "Narx noaniq; ishdan keyin qo‘shimcha summa chiqishi mumkin.",
            "Pul oldindan beriladi — ish bajarilmasa himoya yo‘q.",
            "Firma buyurtmalarni Excel/daftar bilan chalkashtiradi; jadval urishib ketadi.",
            "Shikoyat va nizo uchun markaziy yo‘l yo‘q.",
        ]
    )
    h2("2.2. E-MAKON nima beradi")
    bullets(
        [
            ("Mijozga: ", "tasdiqlangan firmalar, aniq narx, escrow (pul himoyasi), jonli holat, chat."),
            ("Firmaga: ", "mijoz oqimi, buyurtma/jadval/xodim bitta panelda, moliya ko‘rinishi."),
            ("Jamiyatga: ", "yashil makon sifatliroq parvarishlanadi; ish o‘rinlari va shaffoflik."),
            ("Yaratuvchilarga: ", "komissiya/obuna/banner asosida o‘sadigan raqamli biznes."),
        ]
    )
    h2("2.3. Ahamiyati (qisqa)")
    table(
        ["Yo‘nalish", "Nima uchun muhim"],
        [
            ["Iqtisodiy", "Yangi daromad (komissiya, obuna), firmalar uchun mijoz oqimi"],
            ["Ijtimoiy", "Ishonchli xizmat, barqarorroq mehnat, kamroq firibgarlik"],
            ["Shahar / ekologiya", "Bog‘-hovli parvarishi madaniyati va sifati oshadi"],
            ["Texnologik", "O‘zbekistonda niche marketplace namunasi"],
            ["Strategik", "Viloyatlarga kengayish, B2B (mahalla, maktab, korxona)"],
        ],
        widths=[4.5, 12.5],
    )

    # ------------------------------------------------------------------ 3
    h1("3. Tizim qanday ishlaydi?")
    bullets(
        [
            "Mijoz xizmat tanlaydi → firma, sana, slot, manzil, rasm.",
            "Buyurtma yaratiladi (Yangi) → to‘lov → pul escrow’da.",
            "Firma qabul qiladi, xodim va vaqt belgilaydi.",
            "Ish bosqichlari: Qabul → Yo‘lda → Yetib keldi → Ishlayapti → Tugadi.",
            "Platforma pulni firmaga o‘tkazadi (komissiya ushlanadi).",
            "Ixtiyoriy: ball, sharh, chat, yillik parvarish shartnomasi.",
        ]
    )

    # ------------------------------------------------------------------ 4
    h1("4. Qaysi tillar va texnologiyalar bilan qilindi?")
    para(
        "Loyiha bir nechta dasturlash tili va freymvorklardan iborat monorepo. "
        "Foydalanuvchi ko‘radigan matnlar — o‘zbek (lotin); kod — quyidagi tillarda."
    )
    table(
        ["Qism", "Til / texnologiya", "Nima uchun"],
        [
            ["Backend API", "Python (Django, Django REST Framework)", "Biznes qoidalar, ma’lumotlar bazasi, xavfsizlik"],
            ["Ma’lumotlar bazasi", "PostgreSQL (prod), SQLite (dev)", "Buyurtma, foydalanuvchi, firma ma’lumotlari"],
            ["Kesh / navbat", "Redis", "Bildirishnoma navbati, bot FSM"],
            ["Mobil ilova", "Dart (Flutter)", "Android, iOS va Web uchun bitta kod"],
            ["Firma paneli", "TypeScript / JavaScript (Next.js, React)", "Firma operatsion web-panel"],
            ["Superadmin paneli", "TypeScript / JavaScript (Next.js, React)", "Platforma boshqaruvi"],
            ["UI stillar", "CSS / Tailwind CSS", "Dizayn tokenlari, responsive"],
            ["Telegram bot", "Python (aiogram 3)", "Telegram orqali mijoz kanali"],
            ["Deploy", "Docker, nginx, bash/PowerShell", "Server, HTTPS, panellarni chiqarish"],
            ["Interfeys tili", "O‘zbek (lotin yozuvida)", "Barcha ekran matnlari"],
        ],
        widths=[4.0, 7.0, 6.0],
    )
    para(
        "Qisqa: backend va bot — Python; mobil — Dart/Flutter; ikkala admin panel — TypeScript/Next.js; "
        "foydalanuvchi tili — o‘zbek.",
        italic=True,
    )

    # ------------------------------------------------------------------ 5
    h1("5. Mobil ilova — sahifalar va skrinshotlar")
    para(
        "Mijoz asosan shu ilovadan buyurtma beradi, to‘laydi va holatni kuzatadi. "
        "Til: o‘zbek. Texnologiya: Flutter (Dart)."
    )

    h2("5.1. Skrinshotlar")
    add_image_pair(
        SHOTS / "m_onboarding.png",
        SHOTS / "m_login.png",
        caption="2-rasm. Onboarding va kirish (telefon)",
    )
    add_image_pair(
        SHOTS / "m_otp.png",
        SHOTS / "m_home.png",
        caption="3-rasm. SMS kod (OTP) va bosh sahifa",
    )
    add_image_pair(
        SHOTS / "m_service.png",
        SHOTS / "m_slots.png",
        caption="4-rasm. Xizmat tanlash va bo‘sh vaqt (slot)",
    )
    add_image(
        SHOTS / "m_orders.png",
        width_in=3.2,
        caption="5-rasm. Buyurtmalar ro‘yxati — holatni kuzatish",
    )

    h2("5.2. Har bir ekran nima uchun kerak")
    table(
        ["Sahifa / ekran", "Nima uchun kerak", "Foydalanuvchi nima qiladi"],
        [
            ["Splash", "Brend va sessiyani tekshirish", "Avtomatik o‘tish"],
            ["Onboarding", "Mahsulotni tushuntirish", "3 slayd — Keyingi / O‘tkazib yuborish"],
            ["Ro‘yxat / Kirish", "Akkaunt ochish yoki kirish", "Telefon raqam"],
            ["OTP", "Telefonni tasdiqlash", "SMS kodni kiritish"],
            ["Bosh sahifa", "Xizmat topish markazi", "Banner, qidiruv, katalog, firmalar, ball"],
            ["Buyurtmalar", "O‘z buyurtmalarini kuzatish", "Faol / tugallangan / bekor"],
            ["Xabarlar", "Firma va tizim bilan aloqa", "Chat, bildirishnomalar"],
            ["Profil", "Shaxsiy ma’lumot", "Ism, manzil, avatar, chiqish"],
            ["Xizmat tafsiloti", "Xizmat haqida batafsil", "Buyurtma berish"],
            ["Buyurtma oqimi", "Buyurtmani yaratish", "Firma → manzil/vaqt/rasm → tasdiq"],
            ["Maydon skaneri", "Maydonni taxminiy o‘lchash", "m² kiritish"],
            ["To‘lov", "Pul to‘lash", "Click / Payme / test"],
            ["Muvaffaqiyat / batafsil", "Natija va kuzatuv", "Bosqichlar, bekor, chat"],
            ["Chat / Ballar", "Aloqa va loyalty", "Yozishish, mukofot"],
        ],
        widths=[4.5, 5.5, 7.0],
    )

    # ------------------------------------------------------------------ 6
    h1("6. Firma admin paneli — sahifalar va skrinshotlar")
    para(
        "Xizmat ko‘rsatuvchi firma uchun. Login: telefon + parol. "
        "Til: o‘zbek. Texnologiya: Next.js (TypeScript/React)."
    )
    para(bold_prefix="Manzil: ", text="dev — localhost:3002 · prod — admin.e-makon.uz")

    h2("6.1. Skrinshotlar")
    add_image_pair(
        SHOTS / "fa_dashboard.png",
        SHOTS / "fa_schedule.png",
        caption="6-rasm. Firma: bosh sahifa va buyurtma jadvali (Schedule)",
    )
    add_image_pair(
        SHOTS / "fa_staff.png",
        SHOTS / "fa_finance.png",
        caption="7-rasm. Firma: xodimlar va moliya",
    )

    h2("6.2. Har bir sahifa nima uchun kerak")
    table(
        ["Sahifa", "Nima uchun kerak", "Nima qilinadi"],
        [
            ["Login", "Firma adminini tizimga kiritish", "Telefon + parol"],
            ["Bosh sahifa", "Kunlik ko‘rsatkichlar", "Buyurtma, daromad, ogohlantirish, hisobot"],
            ["Buyurtmalar", "Operatsiya markazi", "Qabul, xodim, narx, bosqich, jadval, soxta belgilash"],
            ["Xizmatlar", "Takliflarni boshqarish", "Qo‘shish, narx, rasm, faol/nofaol"],
            ["Xodimlar", "Brigadani yuritish", "Qo‘shish, lavozim, holat, kartochka"],
            ["Mijozlar", "Mijozlar bazasi", "Qidiruv, tarix, chat, Excel"],
            ["Parvarish", "Yillik shartnomalar", "Yaratish, tashrif, tasdiqqa yuborish"],
            ["Moliya", "Pul oqimi", "Escrow, ulush, filtr"],
            ["Xarita", "Joylashtirish", "Yandex: buyurtma/xodim nuqtalari"],
            ["Murojaatlar", "Mijoz bilan yozishish", "Chat, javob, o‘qilmaganlar"],
            ["Hisobotlar", "Davr tahlili", "Statistika, PDF/Excel"],
            ["Firma profili", "Firma ma’lumoti", "Manzil, ish vaqti, ijtimoiy tarmoq"],
            ["Sozlamalar", "Akkaunt sozlamalari", "Profil va boshqa shaxsiy sozlamalar"],
        ],
        widths=[4.0, 5.5, 7.5],
    )

    # ------------------------------------------------------------------ 7
    h1("7. Superadmin paneli — sahifalar va skrinshotlar")
    para(
        "Platforma egasi / E-MAKON markazi uchun. Barcha firmalar, to‘lovlar va sifat nazorati. "
        "Til: o‘zbek. Texnologiya: Next.js (TypeScript/React)."
    )
    para(bold_prefix="Manzil: ", text="dev — localhost:3000 · prod — superadmin.e-makon.uz")

    h2("7.1. Skrinshotlar")
    add_image_pair(
        SHOTS / "sa_dashboard.png",
        SHOTS / "sa_firms.png",
        caption="8-rasm. Superadmin: bosh sahifa va firmalar",
    )
    add_image_pair(
        SHOTS / "sa_schedule.png",
        SHOTS / "sa_reports.png",
        caption="9-rasm. Superadmin: jadval/buyurtmalar va hisobotlar",
    )
    add_image(
        SHOTS / "sa_map.png",
        width_in=5.5,
        caption="10-rasm. Superadmin: xarita / geografik monitoring",
    )

    h2("7.2. Har bir sahifa nima uchun kerak")
    table(
        ["Sahifa", "Nima uchun kerak", "Nima qilinadi"],
        [
            ["Login", "Platforma adminini kiritish", "Telefon + parol"],
            ["Bosh sahifa", "Butun platforma pulsini ko‘rish", "GMV, buyurtma, firma, ogohlantirish"],
            ["Hisobotlar", "Chuqur tahlil", "Davr, jadval, PDF/Excel"],
            ["Firmalar", "Hamkorlarni boshqarish", "Status, sinov, blok, batafsil"],
            ["Firma batafsil", "Bitta firmani chuqur ko‘rish", "Ledger, jazo, savdo taqiqi"],
            ["Foydalanuvchilar", "Mijozlar bazasi", "Qidiruv, bloklash"],
            ["Investorlar", "Investor hamkorlar", "Qo‘shish, shartnoma"],
            ["Xodimlar", "Barcha firmalar xodimlari", "Qidiruv, holat"],
            ["Xizmatlar", "Katalog + moderatsiya", "Tasdiq/rad, katalog turlari"],
            ["Buyurtmalar", "Global buyurtma + moliya", "To‘lov tasdiq, escrow, nizo, jazo"],
            ["Parvarish", "Yillik shartnoma nazorati", "Tasdiq, rad, monitoring"],
            ["Bannerlar", "Ilova reklamalari", "Karusel/promo, muddat"],
            ["Aloqa", "Markaziy aloqa", "Chat monitoring, ticket"],
            ["Ball tizimi", "Loyalty dasturi", "Sozlama, mukofot, tranzaksiya"],
            ["Xarita", "Geografik monitoring", "Firma/xodim/buyurtma joyi"],
            ["Administratorlar", "Firma admin akkauntlari", "Firma + vakolat + parol"],
            ["Sozlamalar", "Shaxsiy profil", "Profil saqlash, tezkor havolalar"],
        ],
        widths=[4.2, 5.3, 7.5],
    )

    # ------------------------------------------------------------------ 8
    h1("8. Uch qismning farqi")
    table(
        ["Savol", "Mobil ilova", "Firma paneli", "Superadmin"],
        [
            ["Kim uchun?", "Mijoz", "Xizmat firmasi", "Platforma egasi"],
            ["Asosiy maqsad", "Buyurtma berish", "Buyurtmani bajarish", "Nazorat va moliya"],
            ["Ma’lumot doirasi", "O‘zining", "O‘z firmasi", "Butun tizim"],
            ["Dasturlash tili", "Dart (Flutter)", "TypeScript (Next.js)", "TypeScript (Next.js)"],
            ["Interfeys tili", "O‘zbek", "O‘zbek", "O‘zbek"],
            ["Kirish", "SMS OTP", "Telefon + parol", "Telefon + parol"],
        ],
        widths=[3.5, 4.5, 4.5, 4.5],
    )

    # ------------------------------------------------------------------ 9
    h1("9. Oldinda qilishimiz kerak bo‘lgan ishlar")
    para(
        "Quyidagi ro‘yxat production’ga chiqish va mahsulotni to‘ldirish uchun prioritet "
        "bo‘yicha tartiblangan. Bu — jamoa uchun yo‘l xaritasi."
    )

    h2("9.1. KRITIK — ishga tushirishdan oldin")
    table(
        ["#", "Ish", "Nima uchun"],
        [
            ["1", "SMS provayder (Eskiz / Playmobile)", "Real mijoz OTP orqali kira olishi uchun"],
            ["2", "Click va Payme to‘liq (callback + imzo)", "To‘lov avtomatik tasdiqlanishi uchun"],
            ["3", "PAYMENTS_TEST_MODE production’da o‘chirish", "Pulsiz «to‘langan» buyurtma bo‘lmasin"],
            ["4", "HTTPS, zaxira, Sentry / logging", "Xavfsiz va barqaror server"],
            ["5", "Mobil Store tayyorligi (imzo, privacy, akkaunt o‘chirish)", "Google Play / App Store"],
            ["6", "Push-bildirishnomalar (FCM)", "Ilova yopiq bo‘lsa ham xabar"],
        ],
        widths=[1.2, 8.5, 7.3],
    )

    h2("9.2. YUQORI — birinchi oy")
    table(
        ["#", "Ish", "Nima uchun"],
        [
            ["1", "Xodim (usta) mobil rejimi", "Brigada telefondan bosqichni belgilaydi"],
            ["2", "Bandlik = firma ish vaqti (work_hours)", "Noto‘g‘ri slotlar chiqmasin"],
            ["3", "Celery/cron: eslatma, ball muddati, sinov tugashi", "Avtomatik biznes qoidalar"],
            ["4", "Bot: slot + to‘lov to‘liq oqim", "Telegram’dan to‘liq buyurtma"],
            ["5", "UI/UX: nav clearance, mobil qidiruv, xarita bir xilligi", "Foydalanish qulayligi"],
            ["6", "Panel imkoniyat bayroqlari UI’da", "Ruxsatsiz menyu chalkashmasin"],
        ],
        widths=[1.2, 8.5, 7.3],
    )

    h2("9.3. O‘RTA — 1–2 oy")
    bullets(
        [
            "Manzil: xaritadan nuqta + reverse geocoding.",
            "Parvarish eslatmalari va tashrif bildirishnomalari.",
            "Shared UI kit (ikki panel dublikatini kamaytirish).",
            "Light mode yoki yuqori kontrast.",
            "Frontend avtomatik testlar.",
            "i18n (rus / kirill) — ixtiyoriy kengayish.",
        ]
    )

    h2("9.4. KEYINGI (o'sish)")
    future = ASSETS / "future.jpg"
    if future.exists():
        add_image(future, width_in=5.2, caption="11-rasm. Kelajak: kengayish va yangi imkoniyatlar")
    bullets(
        [
            "WebSocket — chat va buyurtma real vaqt.",
            "Xodimni xaritada jonli kuzatish.",
            "AI: rasm/maydon bo‘yicha narx taxmini.",
            "B2B: mahalla, maktab, korxona yillik shartnomalar.",
            "Viloyat markazlariga kengayish.",
        ]
    )

    # ------------------------------------------------------------------ 10
    h1("10. Xulosa")
    para(
        "E-MAKON — mijoz, firma va platformani bog‘lovchi to‘liq tizim. "
        "U ishonchsiz xizmat bozorini raqamli, shaffof va himoyalangan jarayonga aylantiradi. "
        "Kod Python, Dart va TypeScript da; interfeys o‘zbek tilida. "
        "Keyingi eng muhim qadamlar: SMS, real to‘lov, Store va push."
    )
    bullets(
        [
            "Nima uchun kerak: ishonch, tartib, pul himoyasi, ish o‘rinlari.",
            "Qanday ishlaydi: buyurtma → escrow → firma → bosqichlar → to‘lov chiqarish.",
            "Kim nima qiladi: mobil — mijoz; admin — firma; superadmin — nazorat.",
        ]
    )

    try:
        doc.save(OUT)
        print(f"OK: {OUT}")
    except PermissionError:
        doc.save(OUT_FALLBACK)
        print(f"Asosiy fayl ochiq — saqlandi: {OUT_FALLBACK}")
        print(f"(Eski faylni yoping va qayta ishga tushiring: {OUT.name})")


if __name__ == "__main__":
    build()
