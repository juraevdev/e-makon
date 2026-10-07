"""E-MAKON texnik topshirig'ini (TZ) Word formatida yig'adi.

Ishga tushirish:  python docs/tz/build_tz.py   (python-docx kerak)
Natija:           docs/E-MAKON_Texnik_topshiriq.docx
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "E-MAKON_Texnik_topshiriq.docx"

GREEN = RGBColor(0x1E, 0x6B, 0x34)
DARK = RGBColor(0x1F, 0x23, 0x21)
GREY = RGBColor(0x5F, 0x66, 0x62)
RED = RGBColor(0xB3, 0x26, 0x1E)
HEADER_FILL = "1E6B34"
ZEBRA_FILL = "EEF5EF"

doc = Document()


# ---------------------------------------------------------------- styles
def setup_styles() -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in (("Heading 1", 16, GREEN), ("Heading 2", 13, GREEN), ("Heading 3", 11.5, DARK)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
    for sec in doc.sections:
        sec.page_height = Cm(29.7)
        sec.page_width = Cm(21.0)
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(1.8)
        sec.bottom_margin = Cm(1.8)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number_footer() -> None:
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("E-MAKON · Texnik topshiriq · sahifa ")
        run.font.size = Pt(8.5)
        run.font.color.rgb = GREY
        r = p.add_run()
        r.font.size = Pt(8.5)
        r.font.color.rgb = GREY
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            r._r.append(el)


# ---------------------------------------------------------------- helpers
def h1(text: str) -> None:
    doc.add_heading(text, level=1)


def h2(text: str) -> None:
    doc.add_heading(text, level=2)


def h3(text: str) -> None:
    doc.add_heading(text, level=3)


def para(text: str, *, bold_prefix: str | None = None, italic: bool = False, color=None, size=None) -> None:
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    r = p.add_run(text)
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)


def bullets(items, style: str = "List Bullet") -> None:
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(2)
        if isinstance(item, tuple):
            r = p.add_run(item[0])
            r.bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)


def numbered(items) -> None:
    bullets(items, style="List Number")


def table(headers, rows, widths=None, font_size: float = 9) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, HEADER_FILL)
    for idx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            text = str(val)
            if text.startswith("**") and "**" in text[2:]:
                end = text.index("**", 2)
                rb = p.add_run(text[2:end])
                rb.bold = True
                rb.font.size = Pt(font_size)
                text = text[end + 2 :]
            r = p.add_run(text)
            r.font.size = Pt(font_size)
            if idx % 2 == 1:
                shade(cells[i], ZEBRA_FILL)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break() -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ================================================================= CONTENT
setup_styles()

# ---------------------------------------------------------------- titul
for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("E-MAKON")
r.bold = True
r.font.size = Pt(40)
r.font.color.rgb = GREEN
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Bog' va landshaft xizmatlari uchun raqamli platforma")
r.font.size = Pt(14)
r.font.color.rgb = GREY
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("TEXNIK TOPSHIRIQ (TZ)")
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = DARK
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Tizimning to'liq tavsifi · hozirgi holat · muammolar · rivojlanish rejasi")
r.font.size = Pt(12)
r.font.color.rgb = GREY
for _ in range(8):
    doc.add_paragraph()
table(
    ["Parametr", "Qiymat"],
    [
        ["Hujjat", "Texnik topshiriq va tizim pasporti"],
        ["Versiya", "1.0"],
        ["Sana", date(2026, 10, 7).strftime("%d.%m.%Y")],
        ["Repozitoriy", "github.com/juraevdev/e-makon (master)"],
        ["Holat", "MVP ishlaydi (lokal demo); production uchun tayyor emas — 10-bo'limga qarang"],
        ["Til", "O'zbek (lotin)"],
    ],
    widths=[4.5, 12.5],
    font_size=10,
)
page_break()

# ---------------------------------------------------------------- mundarija
h1("Mundarija")
for line in [
    "1. Umumiy ma'lumot",
    "2. Atamalar va qisqartmalar",
    "3. Tizim arxitekturasi va texnologiyalar",
    "4. Foydalanuvchi rollari va huquqlar",
    "5. Funksional tavsif (backend modullari, biznes-qoidalar)",
    "6. Mobil ilova (mijoz)",
    "7. Firma admin paneli",
    "8. Superadmin paneli",
    "9. Telegram bot",
    "10. Hozirgi muammolar va kamchiliklar",
    "11. Bajarilgan ishlar",
    "12. Qilinishi kerak bo'lgan ishlar (yo'l xaritasi)",
    "13. Nofunksional talablar",
    "14. Infratuzilma, sozlamalar va ishga tushirish",
    "15. Qabul qilish mezonlari",
    "Ilova A. API endpointlar ro'yxati",
    "Ilova B. Ma'lumotlar modeli",
    "Ilova C. Demo kirish ma'lumotlari va buyruqlar",
]:
    para(line)
page_break()

# ================================================================= 1
h1("1. Umumiy ma'lumot")
h2("1.1. Loyiha nomi va maqsadi")
para(
    "E-MAKON — bog' va landshaft xizmatlarini (gazon parvarishi, daraxt kesish, dorilash, sug'orish tizimi, "
    "landshaft dizayn va h.k.) onlayn buyurtma qilish, bajarish va nazorat qilish uchun yagona raqamli platforma. "
    "Platforma mijozni, xizmat ko'rsatuvchi firmani (hamkorni) va platforma ma'muriyatini bitta tizimda bog'laydi."
)
para("Asosiy maqsadlar:", bold_prefix=None)
bullets(
    [
        "Mijoz ishonchli firmani topishi, narxni oldindan bilishi va qulay vaqtni tanlashi.",
        "To'lov xavfsizligi: pul ish tasdiqlanguncha platforma hisobida ushlanadi (escrow).",
        "Firmalar uchun buyurtma, xodim, jadval, moliya va hisobotlarni bitta panelda yuritish.",
        "Platforma uchun firmalar ustidan nazorat: moderatsiya, komissiya, obuna, jarima, savdo taqiqi.",
    ]
)
h2("1.2. Hal qilinadigan muammo")
table(
    ["Muammo (bozorda)", "E-MAKON yechimi"],
    [
        ["Ishonchli usta topish qiyin (tanish-bilish, Telegram guruhlar)", "Moderatsiyadan o'tgan firmalar, reyting va faqat haqiqiy mijoz qoldiradigan sharhlar"],
        ["Narx noaniq, ish oxirida oshib ketadi", "Firma qat'iy narxni oldindan belgilaydi, mijoz bir nechta firmani solishtiradi"],
        ["Vaqt kelishilmaydi, usta kechikadi", "Bandlik jadvali: bo'sh soat yashil, band soat qizil; server ikki marta band qilishga yo'l qo'ymaydi"],
        ["Oldindan to'lov — qaytarib olish imkonsiz", "Escrow: pul platformada ushlanadi, ish tasdiqlangach firmaga o'tadi yoki mijozga qaytariladi"],
        ["Firmalar buyurtmani daftarda yuritadi", "Firma paneli: buyurtmalar, ish bosqichlari, xodimlar, moliya, PDF hisobot"],
    ],
    widths=[7.5, 9.5],
)
h2("1.3. Tizim tarkibi (qisqacha)")
table(
    ["Qism", "Kim uchun", "Texnologiya", "Papka / port"],
    [
        ["Backend API", "Barcha mijoz ilovalar uchun yagona manba", "Django 5 + DRF + SimpleJWT", "config/, apps/ · :8000"],
        ["Mobil ilova", "Mijoz", "Flutter (Android, iOS, Web)", "mobile/"],
        ["Firma admin paneli", "Hamkor firma admini", "Next.js 16 + React 19 + Tailwind 4", "admin_panel/ · :3002"],
        ["Superadmin paneli", "Platforma egasi", "Next.js 16 + React 19 + Tailwind 4", "superadmin_panel/ · :3000"],
        ["Telegram bot", "Mijoz (qo'shimcha kanal)", "aiogram 3 (faqat HTTP orqali API bilan)", "bot/"],
        ["Investor taqdimoti", "Investorlar", "python-pptx, 16 slayd", "pitch/"],
    ],
    widths=[3.3, 4.2, 5.0, 4.5],
)

# ================================================================= 2
h1("2. Atamalar va qisqartmalar")
table(
    ["Atama", "Ma'nosi"],
    [
        ["Firma / hamkor", "Xizmat ko'rsatuvchi tashkilot (Organization). Har bir firma alohida tenant."],
        ["Default tashkilot", "slug=default — platformaning o'zi; OTP orqali ro'yxatdan o'tgan mijozlar shunga biriktiriladi."],
        ["Katalog turi", "Platforma belgilagan ildiz xizmat (9 ta): masalan, «Gazon parvarish»."],
        ["Firma taklifi (oferta)", "Firmaning katalog turiga bog'langan o'z xizmati va qat'iy narxi; superadmin moderatsiyasidan o'tadi."],
        ["Escrow", "Mijoz to'lovini ish tasdiqlanguncha platforma hisobida ushlab turish mexanizmi."],
        ["Ledger", "Ikki tomonlama buxgalteriya yozuvlari (debet/kredit hisoblar)."],
        ["Slot", "Bandlik jadvalidagi 1 soatlik vaqt oralig'i (08:00–20:00)."],
        ["Sig'im (capacity)", "Firma bir vaqtda bajara oladigan buyurtmalar soni = faol xodimlar soni (kamida 1)."],
        ["Ish bosqichi", "accepted → on_the_way → arrived → working → finished."],
        ["OTP", "SMS orqali yuboriladigan bir martalik 6 xonali kod."],
        ["JWT", "Kirish tokeni (access 12 soat, refresh 30 kun)."],
        ["Parvarish shartnomasi", "Uzoq muddatli (haftalik/oylik) tashriflar shartnomasi (CareContract)."],
        ["MVP", "Ishlaydigan birinchi versiya."],
    ],
    widths=[4.0, 13.0],
)

# ================================================================= 3
h1("3. Tizim arxitekturasi va texnologiyalar")
h2("3.1. Umumiy sxema")
para(
    "Backend API — yagona «haqiqat manbai». Mobil ilova, ikki web panel va Telegram bot faqat REST API orqali ishlaydi "
    "(bot ham ma'lumotlar bazasiga to'g'ridan-to'g'ri ulanmaydi). Bildirishnomalar Redis navbati orqali bot jarayoniga uzatiladi."
)
table(
    ["Oqim", "Tavsif"],
    [
        ["Mobil ilova → API", "HTTPS/HTTP JSON va multipart (rasm), Bearer JWT, 401 da refresh"],
        ["Panellar → API", "fetch, JWT localStorage'da, 401 da bir marta refresh"],
        ["Bot → API", "httpx; X-Bot-Service-Key (Telegram bog'lash), Idempotency-Key (buyurtma)"],
        ["API → Redis → Bot", "Hodisalar (order.created, order.status_changed, order.cancelled, support.reply) RPUSH/BLPOP"],
        ["API → Click/Payme", "Faqat to'lov havolasi yasaladi (callback yo'q)"],
    ],
    widths=[4.0, 13.0],
)
h2("3.2. Backend qatlamlari")
bullets(
    [
        "View → Serializer → Service → Model. Biznes-mantiq services.py fayllarida.",
        "Javob formati: muvaffaqiyat {success, message, data}; xato {success:false, message, errors, code}.",
        "Pagination: 20 ta, ?page_size= bilan 100 tagacha; {count, next, previous, results}.",
        "Xato klasslari: AppError (400), ConflictError (409), ForbiddenError (403), OTPError (400, otp_error).",
        "Telefon normalizatsiyasi: +998XXXXXXXXX.",
        "Swagger: /api/docs/, OpenAPI sxema: /api/schema/, sog'liq: /health/.",
    ]
)
h2("3.3. Texnologiyalar steki")
table(
    ["Qatlam", "Texnologiya va versiya"],
    [
        ["Backend", "Python 3.12 (Docker) / 3.14 (lokal), Django 5.2, DRF 3.18, SimpleJWT 5.5, drf-spectacular, django-filter, django-cors-headers, Pillow"],
        ["Ma'lumotlar bazasi", "Lokal: SQLite (USE_SQLITE=True). Production: PostgreSQL (psycopg 3)"],
        ["Navbat / kesh", "Redis 7 (bildirishnomalar navbati, bot FSM, rate limit). Celery yo'q"],
        ["Mobil", "Flutter 3.44, Dart ^3.9; provider, go_router, http, flutter_secure_storage, image_picker, geolocator, flutter_map, camera, flutter_local_notifications"],
        ["Web panellar", "Next.js 16.3, React 19.3, Tailwind CSS 4, TypeScript 5. Qo'shimcha kutubxona yo'q (grafiklar qo'lda SVG)"],
        ["Bot", "aiogram 3 (lokal 3.31), httpx, cryptography (Fernet), Redis FSM / Memory"],
        ["Deploy", "Dockerfile (gunicorn), docker-compose (redis, api, bot), run.ps1 (Windows lokal)"],
    ],
    widths=[3.5, 13.5],
)

# ================================================================= 4
h1("4. Foydalanuvchi rollari va huquqlar")
table(
    ["Rol", "Qanday kiradi", "Nima qila oladi"],
    [
        ["Mijoz (customer)", "Telefon + SMS kod (OTP) — mobil ilova yoki Telegram bot", "Katalog, firma tanlash, buyurtma, to'lov, kuzatish, bekor qilish, chat, murojaat, sharh, bonus"],
        ["Firma admini (admin)", "Telefon + parol — firma paneli (:3002)", "O'z firmasi: buyurtmalar, bosqichlar, jadval, xodimlar, xizmatlar, mijozlar, parvarish, moliya (ko'rish), chat, hisobot"],
        ["Superadmin", "Telefon + parol — superadmin paneli (:3000)", "Hammasi: firmalar, moderatsiya, escrow amallari, jarima/taqiq, adminlar, bannerlar, ball tizimi, investorlar"],
        ["Xodim (worker)", "Hozircha KIRA OLMAYDI", "Faqat firma admini tomonidan buyurtmaga biriktiriladi (10-bo'limga qarang)"],
        ["Bot servisi", "X-Bot-Service-Key sarlavhasi", "Mijozning Telegram akkauntini bog'lash/uzish"],
    ],
    widths=[3.2, 5.3, 8.5],
)
para("Firma admini uchun qo'shimcha imkoniyat bayroqlari (AdminProfile): can_manage_orders, can_manage_staff, can_view_analytics. "
     "Backend ularni tekshiradi; panel interfeysi hozircha tekshirmaydi.")
para("Multi-tenancy: firma admini faqat o'z tashkiloti ma'lumotlarini ko'radi (OrganizationQuerysetMixin); superadmin hammasini ko'radi.")

# ================================================================= 5
h1("5. Funksional tavsif (backend modullari)")

h2("5.1. Autentifikatsiya (apps/accounts)")
table(
    ["Parametr", "Qiymat"],
    [
        ["OTP kod uzunligi", "6 raqam (secrets.randbelow)"],
        ["OTP amal qilish muddati", "300 soniya"],
        ["Kodni kiritish urinishlari", "5 ta (atomik hisoblagich, hmac.compare_digest)"],
        ["Kod so'rash limiti", "telefon uchun 600 soniyada 5 ta"],
        ["Debug rejim", "OTP_DEBUG_RETURN_CODE=True bo'lsa javobda debug_code qaytadi (lokal demo uchun yoqilgan)"],
        ["Access token", "12 soat"],
        ["Refresh token", "30 kun, rotatsiya yoqilgan, blacklist yo'q"],
        ["Admin login", "POST /auth/admin/login/ — faqat admin va superadmin rollari"],
        ["Kim OTP oladi", "Faqat mijozlar; yangi raqam bo'lsa avtomatik mijoz yaratiladi (default tashkilotga)"],
    ],
    widths=[5.5, 11.5],
)

h2("5.2. Katalog, firma takliflari va moderatsiya (apps/catalog)")
bullets(
    [
        "9 ta katalog turi (seed_catalog): Bepul maslahat, Landshaft dizayn, Daraxt parvarishi, Gazon parvarish, Archaga shakl, "
        "Hasharotlarga qarshi dorilash, Ozuqalash, 1 yillik kafolat, Sug'orish tizimi.",
        "Firma o'z xizmatini yaratadi: katalog turini tanlaydi (yoki yangi tur taklif qiladi), bitta qat'iy narx kiritadi (price_from = price_to).",
        "Yaratish yoki nom/narx/tavsif/rasm o'zgarishi → status «pending», superadminga bildirishnoma.",
        "Superadmin «approve» (kerak bo'lsa yangi katalog turiga aylantiradi) yoki «reject» (sabab majburiy) qiladi; firmaga xabar boradi.",
        "Ommaga ko'rinadi: faol + tasdiqlangan + firmasi faol bo'lgan xizmatlar.",
        "GET /services/ — katalog turlari va har bir firma narxi (firms[], offers_count, min/max narx).",
        "Bannerlar: draft/scheduled/active/archived, joylashuv home/promo, sanalar oralig'i.",
    ]
)

h2("5.3. Firmalar (apps/organizations)")
table(
    ["Funksiya", "Qoida"],
    [
        ["Statuslar", "active, pending, suspended, ended (faqat active ommaga ko'rinadi va buyurtma qabul qiladi)"],
        ["Komissiya", "Har bir firma uchun 0.1–3%, standart 0.30%"],
        ["Tavsiya etilgan stavka", "Bajarilgan buyurtmalar: <5 → 0.30%; 5–19 → 0.25%; 20–49 → 0.20%; 50–99 → 0.15%; 100+ → 0.10%"],
        ["Obuna", "$9/oy yoki $90/yil × o'rinlar soni (faqat hisoblanadi, billing yo'q)"],
        ["Sinov muddati", "Superadmin kunlar sonini belgilaydi (set_trial)"],
        ["Jarima", "add_fine — unpaid_fines va qarzga qo'shiladi"],
        ["Savdo taqiqi", "ban_sales — standart 7 kun (kamida 1); taqiq davrida yangi buyurtma qabul qilinmaydi"],
        ["Ogohlantirish", "send_message(kind=warning) → warnings_count oshadi"],
        ["Shartnomani tugatish", "end_agreement → status ended"],
        ["Sharhlar", "Faqat shu firmada bajarilgan buyurtmasi bor mijoz; har buyurtmaga bitta; reyting = o'rtacha ball"],
        ["Moderatsiya jurnali", "Har bir superadmin amali FirmModerationLog'ga yoziladi"],
        ["Investorlar", "Investor: summa, ulush %, status active/pending/ended"],
    ],
    widths=[4.5, 12.5],
)

h2("5.4. Buyurtmalar (apps/orders)")
h3("Yaratish")
bullets(
    [
        "POST /orders/ — JSON yoki multipart (rasmlar media maydonida, service_ids[i]).",
        "Idempotency-Key sarlavhasi: bir xil kalit bilan qayta yuborilsa, mavjud buyurtma qaytadi (dublikat yo'q).",
        "Firma faol emas yoki savdo taqiqida bo'lsa — firm_unavailable xatosi.",
        "Vaqt tanlangan bo'lsa, tashkilot qatori qulflanadi (select_for_update) va slot sig'imi tekshiriladi.",
        "Firma (firm_id) tanlansa, katalog turi shu firmaning tasdiqlangan taklifiga almashtiriladi.",
    ]
)
h3("Narx hisoblash (server tomonida)")
para("blocks = clamp(ceil(maydon/10), 1, 500); qadam = (max − min)/40 (max yo'q bo'lsa min × 0.04); "
     "har blok narxi = clamp(round(qadam), 5 000, 500 000); jami = min + (blocks − 1) × har blok, [min, max] oralig'ida. "
     "Bir nechta xizmat narxlari qo'shiladi. Narx 0 bo'lsa to'lov talab qilinmaydi. Firma qat'iy narxida maydon ta'sir qilmaydi.")
h3("Holatlar mashinasi")
table(
    ["Holat", "Keyingi ruxsat etilgan holatlar", "Izoh"],
    [
        ["new", "in_review, cancelled", "Narx > 0 bo'lsa, to'lanmaguncha oldinga o'tmaydi (order_payment_required)"],
        ["in_review", "contacted, cancelled", "Bosqich: accepted"],
        ["contacted", "completed, cancelled, in_review", "Bosqichlar: on_the_way, arrived, working"],
        ["completed", "— (yakuniy)", "Bosqich finished; bonus ball beriladi"],
        ["cancelled", "— (yakuniy)", "Escrow «held» bo'lsa, pul avtomatik qaytariladi"],
    ],
    widths=[3.0, 5.5, 8.5],
)
bullets(
    [
        "Ish bosqichlari faqat oldinga yuradi; set_stage oraliq holatlarni avtomatik o'tkazadi.",
        "completed bo'lganda pul firmaga faqat superadmin yoki tizim yakunlaganda o'tkaziladi (firma o'ziga to'lay olmaydi).",
        "To'lovdan keyin narxni faqat superadmin o'zgartira oladi.",
        "Har o'tishda: tarix (OrderStatusHistory), mijozga ilova ichidagi bildirishnoma, Telegram hodisasi.",
        "Xodim biriktirish: shu firmaning ishdan bo'shatilmagan xodimi.",
    ]
)

h2("5.5. Bandlik jadvali (yangi funksiya, apps/orders/schedule.py)")
table(
    ["Parametr", "Qiymat"],
    [
        ["Ish vaqti", "08:00–20:00, 1 soatlik slotlar"],
        ["Xizmat davomiyligi", "standart 60 daqiqa; firma 30–720 daqiqagacha belgilaydi yoki +30 daq / +1 soat uzaytiradi"],
        ["Sig'im", "faol xodimlar soni (is_active va employment_status=active), kamida 1"],
        ["Bandlikni hisoblash", "faqat new / in_review / contacted holatidagi, vaqti belgilangan buyurtmalar"],
        ["Rang", "free — yashil; busy (bandlar soni ≥ sig'im) — qizil; past (o'tgan) — kulrang"],
        ["Tekshiruv", "slot_outside_hours, slot_in_past, slot_busy xato kodlari"],
        ["Uzaytirish", "hech qachon bloklanmaydi; vaqtni ko'chirish esa sig'imga tekshiriladi"],
        ["API", "GET /partners/{id}/availability/?date= (ommaviy); GET /admin/orders/availability/; POST /admin/orders/{id}/schedule/"],
        ["Ko'rinishi", "mobil ilovada slot tanlash, firma va superadmin panelida «Bandlik jadvali» (ScheduleBoard)"],
    ],
    widths=[4.5, 12.5],
)

h2("5.6. To'lov, escrow va buxgalteriya")
numbered(
    [
        "Mijoz provayderni tanlaydi (click / payme / test) → escrow «awaiting_payment», to'lov «pending», checkout havola.",
        "Mijoz «To'lovni amalga oshirdim» → to'lov «submitted».",
        "Superadmin «mark_paid» → escrow «held», to'lov «confirmed» (yoki «reject_payment»).",
        "Ish yakunlanadi → superadmin «release»: komissiya platformaga, qolgani firmaga.",
        "Bekor qilish/nizo → «refund» (mijozga qaytarish), «dispute», «punish_firm», «punish_user».",
    ]
)
table(
    ["Hodisa", "Ledger yozuvi (debet → kredit)"],
    [
        ["Ushlash (hold)", "customer_cash → platform_escrow"],
        ["O'tkazish (release)", "platform_escrow → platform_revenue (komissiya) va platform_escrow → firm_payable (firma ulushi)"],
        ["Qaytarish (refund)", "platform_escrow → refund_payable"],
        ["Firma jarimasi", "firm_debt → platform_revenue (standart 100 000 so'm)"],
    ],
    widths=[4.0, 13.0],
)
para("Komissiya: fee = summa × stavka / 100; firma ulushi = summa − fee. Stavka: buyurtmadagi → firmaniki → 0.30%.")
para("Test provayder (PAYMENTS_TEST_MODE) to'lovni darhol tasdiqlaydi — faqat demo uchun.", color=RED)

h2("5.7. Xodimlar va adminlar (apps/staff)")
bullets(
    [
        "EmployeeProfile: mutaxassislik (6 tur), reyting, holat active / on_leave / dismissed, lavozim, ishga qabul/bo'shatish sanalari, ko'nikmalar (20 tagacha).",
        "Xodimni o'chirib bo'lmaydi — faqat ishdan bo'shatish (sabab majburiy, user.is_active=False).",
        "Xodim kartasi: buyurtmalar statistikasi, daromad, parvarish tashriflari, oxirgi 10 buyurtma.",
        "AdminProfile: firma admini va uning imkoniyat bayroqlari; superadmin yaratadi.",
    ]
)

h2("5.8. Parvarish shartnomalari (apps/care)")
bullets(
    [
        "Mijoz turi: xonadon / tashkilot. Chastota: haftalik, ikki haftada bir, oylik. To'lov: oylik, choraklik, oldindan.",
        "Holatlar: draft → pending → (approve) active ⇄ paused → completed; rejected; cancelled.",
        "Tasdiqlanganda tashriflar avtomatik yaratiladi (500 tagacha, muddat 5 yilgacha); jami = tashrif narxi × tashriflar soni.",
        "Tashrif holatlari: scheduled, approved, postponed, done, not_done, rejected (reminded va awaiting_report hali ishlatilmaydi).",
        "Firma yaratadi, superadmin tasdiqlaydi/rad etadi; mijoz mobil API'da faqat o'qiydi.",
    ]
)

h2("5.9. Murojaatlar va chat (apps/support)")
bullets(
    [
        "Ticketlar: open / in_progress / resolved / closed; ustuvorlik low / normal / high; ichki (internal) xabarlar mijozga ko'rinmaydi.",
        "Chat: firma ↔ mijoz xonasi (bitta juftlik uchun bitta xona), xabar 4000 belgigacha, oxirgi 200 xabar, ?after= bilan polling.",
        "Superadmin istalgan chatni o'qiydi va «E-Makon» nomidan yozadi.",
        "Firma admini faqat buyurtma/ro'yxatdan o'tish/oldingi xona orqali bog'langan mijoz bilan chat ocha oladi.",
    ]
)

h2("5.10. Bonus ballar (apps/loyalty)")
bullets(
    [
        "Har 10 000 so'm uchun 1 ball (buyurtma yakunlanganda, bir marta).",
        "Almashtirish uchun kamida 50 ball; mukofotlar ro'yxatini superadmin boshqaradi.",
        "Ball muddati 12 oy deb sozlangan, lekin hozircha qo'llanmaydi.",
    ]
)

h2("5.11. Bildirishnomalar (apps/notifications)")
bullets(
    [
        "Ilova ichidagi xabarlar (UserNotification): order, chat, care, system, offer, bonus.",
        "Telegram: Redis navbati orqali bot yuboradi (foydalanuvchi Telegram bog'lagan bo'lsa).",
        "Push (FCM) va SMS bildirishnomalar YO'Q.",
    ]
)

h2("5.12. Analitika, hisobot va xarita (apps/analytics)")
bullets(
    [
        "Dashboard: davr (kun, hafta, oy, mavsum, yil, ixtiyoriy oraliq); mijozlar, firmalar, buyurtmalar, aylanma, platforma ulushi, o'rtacha chek, dinamika, top/kam xizmatlar, firma daromadlari.",
        "Hisobot: GET /admin/reports/bundle/ JSON qaytaradi, PDF brauzerda window.print orqali yasaladi.",
        "Xarita: firmalar, faol xodimlar, ochiq buyurtmalar, parvarish shartnomalari koordinatalari (har biridan 200 tagacha).",
    ]
)

# ================================================================= 6
h1("6. Mobil ilova (mijoz)")
para("Flutter, qorong'i mavzu, faqat o'zbek tili, provider + go_router. API manzili --dart-define=API_BASE_URL orqali beriladi.")
table(
    ["Ekran", "Nima ko'rsatadi / nima qiladi", "Ma'lumot manbai"],
    [
        ["Splash, Onboarding", "Logotip, 3 ta tanishtiruv sahifasi", "Lokal"],
        ["Ro'yxatdan o'tish / Kirish", "Ism, familiya, telefon (+998), «SMS kod olish»", "API (OTP)"],
        ["OTP", "6 katakli kod, 59 soniya taymer, demo rejimida kod avtomatik to'ladi", "API"],
        ["Bosh sahifa", "Salomlashish, qidiruv, karusel, ballar, hamkorlar, xizmatlar to'ri, eng yaqin ofislar xaritasi, takliflar", "API + demo zaxira"],
        ["Hamkor oynasi", "Faoliyat, ishlar galereyasi, sharhlar, aloqa, navigator, chat", "API (sharhlar lokal)"],
        ["Xizmat sahifasi", "Tavsif, narx oralig'i, firmalar takliflari (narx/reyting/masofa bo'yicha saralash)", "API"],
        ["Buyurtma (3 qadam)", "Firma tanlash → manzil, maydon, sana, VAQT SLOTI (yashil/qizil), telefon, izoh, 2–3 rasm → tasdiqlash", "API"],
        ["Maydon skaneri", "GPS bilan burchaklarni yurib o'lchash yoki uzunlik × kenglik", "Qurilma"],
        ["To'lov", "Click / Payme havolasi, «To'ladim», holatni yangilash, test to'lov", "API"],
        ["Buyurtmalar", "Ro'yxat, filtrlar, holat, to'lov, vaqt", "API"],
        ["Buyurtma tafsiloti", "5 bosqichli kuzatuv, ETA, firma va xodim, chat, bekor qilish (20 soniyada yangilanadi)", "API"],
        ["Xabarlar", "Bildirishnomalar va chatlar ro'yxati", "API"],
        ["Chat", "Firma bilan yozishma (4 soniyada polling)", "API"],
        ["Profil", "Avatar, ism, telefon, ballar, sevimlilar, ma'lumotlarni tahrirlash, chiqish", "API + lokal"],
    ],
    widths=[3.5, 9.5, 4.0],
)
para("Faqat lokal (serverga saqlanmaydi): bonus ballar va mukofotlar, sharhlar, sevimlilar, avatar. "
     "Parvarish shartnomalari ekrani yo'q. Karusel va takliflar uchun mobil /home/carousel/ va /offers/ manzillarini so'raydi — "
     "backendda bunday endpointlar yo'q, shuning uchun ular har doim demo ma'lumotdan olinadi.")
para("Platforma sozlamalari: Android applicationId uz.emakon.emakon_app (targetSdk 36), iOS bundle uz.emakon.emakonApp (iOS 13+), "
     "ruxsatlar: internet, joylashuv, kamera, bildirishnoma, galereya. Web versiyasi qo'shilgan (mobile/web).")

# ================================================================= 7
h1("7. Firma admin paneli (admin_panel, :3002)")
table(
    ["Sahifa", "Imkoniyatlar"],
    [
        ["Bosh sahifa", "KPI (faol buyurtmalar, mijozlar, reyting, aylanma), diqqat talab qiladiganlar, dinamika grafigi, top xizmatlar, PDF hisobot"],
        ["Buyurtmalar", "8 ta filtr, Bandlik jadvali, buyurtma oynasi: vaqt va davomiylik, +30 daq/+1 soat, xodim biriktirish, masofa/ETA, keyingi bosqich (to'lovdan keyin), bekor qilish, narxni to'lovgacha o'zgartirish, CSV, Google yo'nalish"],
        ["Xizmatlar", "O'z xizmatlari va moderatsiya holati; yaratish/tahrirlash (narx ≥ 1000, rasm ≤ 5 MB), faollashtirish, o'chirish, CSV"],
        ["Xodimlar", "Faol / ta'tilda / bo'shatilgan; yaratish, tahrirlash, shaxsiy karta, holatni o'zgartirish, CSV"],
        ["Mijozlar", "Segmentlar (yangi, qaytgan, faol, nofaol), mijoz kartasi va buyurtmalar tarixi, chat ochish, CSV"],
        ["Parvarish", "Shartnoma yaratish (chastota, hafta kunlari, xodimlar, narx), qoralama/yuborish, pauza, davom, yakunlash, bekor qilish, tashrif amallari"],
        ["Moliya", "Escrow KPI, buyurtmalar bo'yicha escrow, ledger (faqat o'qish)"],
        ["Xarita", "Google Maps (kalit bo'lsa) yoki embed; firma, buyurtmalar, xodimlar, parvarish qatlamlari"],
        ["Murojaatlar", "Mijoz chatlari (tezkor javoblar), ticketlarga javob, yopish"],
        ["Hisobotlar", "Davr bo'yicha KPI, holatlar, grafik, PDF"],
        ["Firma profili", "Kontaktlar, geolokatsiya, ish vaqti, ijtimoiy tarmoqlar, E-Makon xabarlari, jarimalar, sharhlar"],
        ["Sozlamalar", "Shaxsiy profil, platforma qoidalari"],
    ],
    widths=[3.3, 13.7],
)
para("Ogohlantirish bannerlari: firma bloklangan/tugatilgan, savdo taqiqi, to'lanmagan jarima, qarz. O'qilmagan chatlar soni har 20 soniyada yangilanadi.")

# ================================================================= 8
h1("8. Superadmin paneli (superadmin_panel, :3000)")
table(
    ["Sahifa", "Imkoniyatlar"],
    [
        ["Umumiy ko'rinish", "Firmalar, faol buyurtmalar, investorlar, aylanma, platforma ulushi; dinamika; top xizmatlar; PDF"],
        ["Hisobotlar", "Firma daromadlari jadvali (stavka, tavsiya stavka, ulush), grafiklar, PDF"],
        ["Firmalar", "Kartalar, filtrlar; yangi firma (asosiy, moliya, obuna kalkulyatori, joylashuv, sinov muddati)"],
        ["Firma sahifasi", "KPI, ledger, sharhlar, moderatsiya jurnali; stavka, tavsiya stavka, qarz, sinov, blok/blokdan chiqarish, tugatish, xabar/ogohlantirish, jarima, savdo taqiqi"],
        ["Buyurtmalar", "Firma tanlash + Bandlik jadvali, filtrlar, buyurtma oynasi: escrow ochish, to'lovni tasdiqlash/rad etish, firmaga o'tkazish, qaytarish, nizo, firmani/mijozni jazolash"],
        ["Xizmatlar", "Moderatsiya navbati, firma takliflari, katalog turlari; tasdiqlash (birlashtirish), rad etish, CSV"],
        ["Foydalanuvchilar", "Mijozlar, ballar, bloklash/blokdan chiqarish"],
        ["Xodimlar / Administratorlar", "Yaratish, tahrirlash, faollashtirish; admin imkoniyat bayroqlari"],
        ["Parvarish", "Shartnomalarni tasdiqlash, rad etish, pauza, bekor qilish, izoh"],
        ["Bannerlar", "Yaratish, tahrirlash, o'chirish"],
        ["Aloqa", "Barcha firma-mijoz chatlarini kuzatish va yozish; ticketlar"],
        ["Ball tizimi", "Qoidalar (so'm/ball, minimum, muddat), mukofotlar, oxirgi tranzaksiyalar"],
        ["Investorlar", "Investorlar kartalari, qo'shish, shartnomani tugatish"],
        ["Xarita", "Firmalar, xodimlar, buyurtmalar (soddalashtirilgan proyeksiya)"],
        ["Sozlamalar", "Profil, adminlar"],
    ],
    widths=[3.6, 13.4],
)

# ================================================================= 9
h1("9. Telegram bot (bot/)")
bullets(
    [
        "Faqat HTTP orqali API bilan ishlaydi; sessiyalar bot SQLite'ida, tokenlar Fernet bilan shifrlangan.",
        "Kirish: kontakt/telefon → OTP → Telegram akkauntni bog'lash.",
        "Asosiy menyu: Xizmat buyurtma qilish, Buyurtmalarim, Manzilim, Murojaat, Profil.",
        "Buyurtma: xizmat → savollar (har xizmat uchun kodda yozilgan) → rasm/video (10 MB), joylashuv → xulosa → yuborish (Idempotency-Key).",
        "Buyurtmalar ro'yxati, ko'rish, bekor qilish; murojaatlar yaratish va javob berish; profil va manzil.",
        "Bildirishnomalar: Redis navbatidan (order.created, order.status_changed, order.cancelled, support.reply).",
        "Rejimlar: polling (ishlaydi), webhook (Python 3.14 da ishlamasligi mumkin).",
        "Botda yo'q: vaqt sloti tanlash, to'lov, ballar, chat, parvarish, akkauntni uzish, /help.",
    ]
)

# ================================================================= 10
h1("10. Hozirgi muammolar va kamchiliklar")
para("Daraja: KRITIK — production'ga chiqishdan oldin albatta hal qilinishi kerak; YUQORI — birinchi oyda; O'RTA — reja asosida; PAST — texnik qarz.", italic=True, color=GREY)

h2("10.1. KRITIK")
table(
    ["#", "Muammo", "Qayerda", "Oqibat"],
    [
        ["K1", "SMS provayder ulanmagan: provayder sozlansa ham kod yuborilmaydi (TODO)", "apps/accounts/services/otp.py:45–62", "Real foydalanuvchi tizimga kira olmaydi; hozir faqat demo kod bilan ishlaydi"],
        ["K2", "Click/Payme: faqat havola yasaladi, callback/webhook va imzo tekshiruvi yo'q; to'lov superadmin tomonidan qo'lda tasdiqlanadi", "apps/orders/payments.py:47–78", "Avtomatik to'lov qabul qilinmaydi"],
        ["K3", "PAYMENTS_TEST_MODE base sozlamada True; mijoz «soxta to'lov» tugmasi bilan to'lovni o'zi tasdiqlay oladi", "config/settings/base.py:205, mobile payment_screen.dart", "Pulsiz buyurtma «to'langan» bo'lib qoladi"],
        ["K4", "Xavfli standart sozlamalar: SECRET_KEY='unsafe-dev-secret-key', ALLOWED_HOSTS=['*'], wsgi/asgi standart holda local (DEBUG=True); HSTS/SSL redirect yo'q", "config/settings/base.py:17–19, config/wsgi.py", "Xavfsizlik teshigi, OTP kod javobda chiqishi mumkin"],
        ["K5", "Production infratuzilma yo'q: docker-compose'da PostgreSQL yo'q, nginx/HTTPS yo'q, collectstatic yo'q, media prodda berilmaydi, zaxira nusxa yo'q", "docker-compose.yml, Dockerfile, config/urls.py:21", "Serverga chiqarib bo'lmaydi"],
        ["K6", "Mobil release: debug kalit bilan imzolanadi, API manzili HTTP LAN IP (192.168.0.117), cleartext ruxsat; maxfiylik siyosati va akkauntni o'chirish yo'q", "mobile/android/app/build.gradle.kts, api_config.dart", "Google Play / App Store qabul qilmaydi"],
        ["K7", "XSS: PDF hisobot server ma'lumotini (firma nomi) escape qilmasdan document.write qiladi; tokenlar localStorage'da", "panellar: (admin)/page.tsx, hisobotlar/page.tsx", "Firma nomiga skript yozib superadmin sessiyasini o'g'irlash mumkin"],
        ["K8", "Admin parol bilan kirishda urinishlar cheklovi (throttling) yo'q; JWT blacklist yo'q — chiqishda token bekor qilinmaydi", "apps/accounts/views.py:62–90, base.py:150", "Parolni tanlab topish xavfi"],
    ],
    widths=[0.9, 6.2, 4.6, 5.3],
    font_size=8.5,
)

h2("10.2. YUQORI")
table(
    ["#", "Muammo", "Qayerda"],
    [
        ["Y1", "Xodim (worker) roli tizimga umuman kira olmaydi; ish bosqichlarini faqat firma admini o'zgartiradi", "otp.py:65–71, accounts/views.py:77"],
        ["Y2", "Push-bildirishnomalar (FCM) yo'q — ilova yopiq bo'lsa mijoz xabar olmaydi", "mobile notification_service.dart"],
        ["Y3", "Mobil profil maydonlari backend bilan mos emas: points ↔ loyalty_points, avatar_path ↔ avatar, company va address saqlanmaydi (backendda home_address)", "mobile auth_provider.dart, apps/accounts/serializers.py"],
        ["Y4", "Mobil bonus, sharhlar, sevimlilar, avatar faqat telefonda saqlanadi; backenddagi loyalty va reviews API ishlatilmaydi; parvarish ekrani yo'q", "mobile features/*"],
        ["Y5", "Mobil /home/carousel/ va /offers/ so'raydi — backendda yo'q; har doim demo ma'lumot. Backenddagi /banners/ ishlatilmaydi. API xato bersa soxta firmalar (soxta telefonlar) ko'rinadi", "catalog_provider.dart, demo_content.dart"],
        ["Y6", "Superadmin admin/xodim yaratganda firma tanlab bo'lmaydi → default tashkilotga tushadi; firma yaratishda firma admin akkaunti yaratilmaydi", "superadmin administratorlar, xodimlar, firmalar sahifalari"],
        ["Y7", "Mijoz ish boshlanganidan keyin ham bekor qila oladi va pul avtomatik qaytadi (firma ish qilib qo'ygan bo'lishi mumkin)", "apps/orders/views.py:82–94, services.py:287"],
        ["Y8", "set_trial faol firmani «pending» qiladi — firma ro'yxatdan yo'qoladi va buyurtma qabul qilmaydi", "apps/organizations/services.py:133"],
        ["Y9", "Bandlik firmaning o'z ish vaqtini (work_hours) hisobga olmaydi — hamma uchun 08–20; faqat boshlanish vaqti tekshiriladi", "apps/orders/schedule.py"],
        ["Y10", "~900 qatorlik ish (bandlik jadvali, seed_demo, web) commit qilinmagan; commit xabarlari ma'nosiz («stuff»); CI yo'q", "git"],
    ],
    widths=[0.9, 10.6, 5.5],
    font_size=8.5,
)

h2("10.3. O'RTA")
bullets(
    [
        ("500 xatolar: ", "int() validatsiyasiz (organizations/views.py:86,160; analytics/views.py:26,49); admin ticket yaratishda IntegrityError; xodim telefoni dublikat bo'lsa IntegrityError."),
        ("Bot: ", "BOT_USE_REDIS standart false — bildirishnomalar va limitlar o'chiq; webhook rejimi Python 3.14 da buzilishi mumkin; savollar kodda qotirilgan; vaqt tanlash va to'lov yo'q."),
        ("Biznes-qoidalar chala: ", "ball muddati qo'llanmaydi; mukofot almashtirilganda haqiqiy natija yo'q; obuna to'lovi va sinov muddati nazorati yo'q; jarimani to'lash endpointi yo'q; tashrif eslatmalari yo'q (rejalashtiruvchi/Celery yo'q)."),
        ("Panellar: ", "imkoniyat bayroqlari UI'da tekshirilmaydi (403 xatolar); ro'yxatlar 50–200 ta bilan cheklangan va mijoz tomonida filtrlanadi; superadmin xaritasi haqiqiy xarita emas; firma profilida buzilgan belgilar (mojibake); jazo summasi kodda 100 000 so'm / 7 kun; confirm/prompt oynalari; ba'zi sahifalarda xato ushlanmaydi."),
        ("Mobil: ", "GPS manzil matnga aylanmaydi (reverse geocoding yo'q); rasm faqat galereyadan va 2–3 ta majburiy; chat va buyurtma polling orqali; web'da sahifani yangilasa ba'zi ekranlar ochilmaydi; i18n yo'q; sanalar inglizcha oy nomlari bilan."),
        ("Monitoring: ", "LOGGING, Sentry, metrikalar yo'q; Redis parolsiz; API so'rovlar limiti (throttling) yo'q."),
    ]
)

h2("10.4. PAST (texnik qarz)")
bullets(
    [
        "README fayllari eskirgan (admin_panel README — superadmin nusxasi; .env.example panellarda yo'q; mashinaga xos yo'llar).",
        "PAYMENTS_TEST_MODE .env.example'da yo'q; requirements.txt versiyalari qotirilmagan.",
        "Ikki panel o'rtasida ~20 ta bir xil fayl (umumiy paket yo'q); superadmin «hamkorlar» sahifasi menyuda yo'q (eski dublikat).",
        "Web manifest va Android ikonka shablon holatda; gradle.properties'da dasturchi JDK yo'li.",
        "Frontendlarda avtomatik testlar yo'q; pitch/ (~25 MB) .gitignore'da yo'q.",
    ]
)

# ================================================================= 11
h1("11. Bajarilgan ishlar")
h2("11.1. Platforma bo'yicha (kodda mavjud va ishlaydi)")
table(
    ["Yo'nalish", "Holat"],
    [
        ["Backend API", "11 modul, 80+ endpoint, Swagger hujjatlari, multi-tenancy, rollar va imkoniyat bayroqlari"],
        ["Kirish", "SMS-kod (OTP) oqimi xavfsiz (hash, urinishlar, limit), admin parol bilan kirish, JWT refresh"],
        ["Katalog", "9 xizmat turi, firma takliflari, narx solishtirish, moderatsiya oqimi"],
        ["Buyurtma", "Idempotent yaratish, server narx hisoblashi, rasm yuklash, holatlar mashinasi, 5 ish bosqichi, tarix"],
        ["Moliya", "Escrow (6 holat), ikki tomonlama ledger, komissiya, qaytarish, nizo, jarima, savdo taqiqi"],
        ["Firma boshqaruvi", "Status, komissiya va tavsiya stavka, obuna kalkulyatori, sinov, ogohlantirish, moderatsiya jurnali"],
        ["Parvarish", "Shartnomalar, tashriflar jadvalini avtomatik yaratish, tasdiqlash oqimi"],
        ["Aloqa", "Firma–mijoz chati, superadmin monitoringi, ticketlar"],
        ["Bonus", "Ball yig'ish va mukofotlar (backend)"],
        ["Analitika", "Dashboard, hisobot, xarita, CSV/PDF eksport"],
        ["Mijoz kanallari", "Flutter mobil ilova (Android/iOS/Web) va Telegram bot"],
        ["Panellar", "Firma paneli (12 sahifa) va superadmin paneli (18 sahifa)"],
        ["Testlar", "99 ta avtomatik test (89 backend + 10 bot)"],
    ],
    widths=[3.8, 13.2],
)
h2("11.2. Oxirgi ish sessiyasida (07.10.2026) tuzatilgan va qo'shilgan")
table(
    ["Ish", "Tafsilot"],
    [
        ["«SMS kod olish» xatosi tuzatildi", "Backend ishga tushirildi, Flutter web uchun CORS (config/settings/local.py) ruxsat berildi, mobil ilova tarmoq xatolarini tushunarli xabar bilan ko'rsatadi (api_client.dart). Demo uchun kod ekranda chiqadi"],
        ["Buyurtmada rasm yuklash tuzatildi", "Multipart so'rov to'g'rilandi (service_ids[i], koordinatalar, media fayllar baytlardan), web'da rasm ko'rinishi va ✕ o'chirish tugmasi, xatolar ushlanadi"],
        ["Ilova kamchiliklari", "Flutter analyzer xatolari (motion.dart, home_screen.dart), profil avatari web'da, Flutter web platformasi qo'shildi (mobile/web)"],
        ["Bandlik jadvali (yangi)", "Order: scheduled_start, duration_minutes (migratsiya 0010); schedule.py; ommaviy va admin availability API; vaqtni ko'chirish/uzaytirish API; ikki panelda ScheduleBoard (yashil/qizil/kulrang); mobil slot tanlash; band slotga yozilsa qayta tanlashga qaytaradi"],
        ["Testlar", "7 ta yangi bandlik testi; eski testdagi qotirilgan sana dinamik qilindi; jami 99 test o'tadi"],
        ["Demo ma'lumotlar", "manage.py seed_demo: 2 firma, xodimlar, takliflar, 40+ buyurtma, sharhlar, ertangi bandlik"],
        ["Investor taqdimoti", "pitch/E-MAKON_investor_taqdimoti.pptx va .pdf (16 slayd, skrinshotlar bilan)"],
    ],
    widths=[4.2, 12.8],
)

# ================================================================= 12
h1("12. Qilinishi kerak bo'lgan ishlar (yo'l xaritasi)")
para("Muddatlar taxminiy, 1–2 dasturchi jamoasi uchun. Har bir vazifa uchun qabul mezoni ko'rsatilgan.", italic=True, color=GREY)

h2("Bosqich 0 — Barqarorlashtirish (1 hafta)")
table(
    ["Vazifa", "Qabul mezoni"],
    [
        ["Commit qilinmagan ishlarni ma'noli xabarlar bilan commit qilish, .gitignore (pitch/, _tmp)", "git status toza; har commit nima qilganini aytadi"],
        ["CI (GitHub Actions): backend testlar, flutter analyze, next lint + build", "Har PR'da avtomatik tekshiruv yashil"],
        ["PDF hisobotda HTML escape (XSS)", "Firma nomida <script> bo'lsa ham bajarilmaydi"],
        ["DRF throttling (admin login, OTP — IP bo'yicha) va JWT blacklist + logout endpoint", "10 noto'g'ri urinishdan keyin 429; chiqishdan keyin refresh ishlamaydi"],
        ["PAYMENTS_TEST_MODE standart False; .env.example to'ldirish", "Production'da test-pay 404/403"],
        ["500 xatolarni tuzatish (int validatsiya, IntegrityError)", "Noto'g'ri qiymatlar 400 qaytaradi"],
        ["Firma profilidagi buzilgan belgilarni tuzatish", "Matnlar to'g'ri ko'rinadi"],
    ],
    widths=[9.0, 8.0],
)

h2("Bosqich 1 — Ishga tushirish (2–4 hafta)")
table(
    ["Vazifa", "Qabul mezoni"],
    [
        ["SMS provayder (Eskiz yoki Playmobile) integratsiyasi", "Real raqamga 10 soniya ichida kod keladi; debug kod o'chiq"],
        ["Click va Payme to'liq integratsiyasi: callback (Payme JSON-RPC, Click prepare/complete), imzo tekshiruvi, provider_transaction_id", "Sandbox'da to'lov avtomatik «held» bo'ladi, qo'lda tasdiqlash shart emas"],
        ["Production infratuzilma: PostgreSQL, nginx + HTTPS (Let's Encrypt), gunicorn workerlar, collectstatic, media (S3 yoki volume), kunlik zaxira, Sentry, LOGGING", "Domen orqali HTTPS ishlaydi; zaxiradan tiklash sinovdan o'tgan"],
        ["Production sozlamalar: SECRET_KEY/ALLOWED_HOSTS majburiy, wsgi → production, HSTS, SSL redirect", "manage.py check --deploy ogohlantirishsiz"],
        ["Mobil release: imzolash kaliti, HTTPS API manzili, cleartext o'chirish, ikonka, maxfiylik siyosati va foydalanish shartlari havolasi, akkauntni o'chirish", "Google Play internal test va TestFlight'ga yuklangan"],
        ["Push-bildirishnomalar (FCM): buyurtma holati, chat, to'lov", "Ilova yopiq bo'lsa ham xabar keladi"],
    ],
    widths=[9.0, 8.0],
)

h2("Bosqich 2 — Mahsulotni to'ldirish (1–2 oy)")
table(
    ["Vazifa", "Qabul mezoni"],
    [
        ["Xodim (usta) uchun kirish va mobil rejim: o'z buyurtmalari, «yo'ldaman / yetib keldim / ishlayapman / tugatdim», rasm hisobot", "Bosqichlarni xodim telefondan o'zgartiradi, mijoz real vaqtda ko'radi"],
        ["Mobil ↔ backend maydonlarini moslash (loyalty_points, avatar yuklash, home_address)", "Profil ma'lumotlari qayta kirganda saqlanib qoladi"],
        ["Mobil: bonus, sharhlar, sevimlilar → API; parvarish shartnomalari ekrani", "Ma'lumot boshqa qurilmada ham ko'rinadi"],
        ["Karusel → /banners/, takliflar uchun endpoint; demo zaxirani faqat dev rejimida qoldirish", "Production'da soxta firma ko'rinmaydi"],
        ["Manzil: xaritadan nuqta tanlash va reverse geocoding; rasm kameradan ham", "Manzil matni avtomatik to'ladi"],
        ["Bandlik: firmaning ish vaqti va dam olish kunlari; xodim bo'yicha taqsimlash", "Firma 09–18 ishlasa, 08:00 slot ko'rinmaydi"],
        ["Superadmin: admin/xodim yaratishda firma tanlash; firma yaratishda admin akkaunt", "Yangi firma darhol panelga kira oladi"],
        ["Bekor qilish siyosati: ish boshlangandan keyin bekor qilish — nizo orqali", "«working» holatida avtomatik qaytarish yo'q"],
        ["Rejalashtiruvchi (Celery beat yoki cron): ball muddati, sinov muddati tugashi, parvarish eslatmalari, obuna to'lovi", "Muddati o'tgan ballar avtomatik yechiladi; eslatmalar yuboriladi"],
        ["Bot: vaqt sloti tanlash, to'lov havolasi, akkauntni uzish; Redis standart yoqilgan", "Botdan to'liq buyurtma va to'lov qilinadi"],
        ["Panellar: imkoniyat bayroqlarini UI'da hisobga olish; server tomonida pagination va filtr", "Cheklangan admin ruxsatsiz menyuni ko'rmaydi"],
    ],
    widths=[9.0, 8.0],
)

h2("Bosqich 3 — O'sish (3–6 oy)")
bullets(
    [
        "Real vaqt: WebSocket (chat, buyurtma holati), xodimni xaritada jonli kuzatish.",
        "Ko'p tillilik: o'zbek (kirill), rus tili.",
        "AI: rasm va maydon bo'yicha narxni taxmin qilish.",
        "B2B: korxona, mahalla va maktablar uchun yillik shartnomalar.",
        "Viloyat markazlariga kengayish, hamkorlik va marketing dasturi.",
        "Panellar uchun umumiy UI paketi, frontend testlari, analitika (mahsulot metrikalari).",
    ]
)

# ================================================================= 13
h1("13. Nofunksional talablar")
table(
    ["Talab", "Hozirgi holat", "Maqsad"],
    [
        ["Xavfsizlik", "OTP hash, urinishlar limiti, tenant izolyatsiyasi; throttling va blacklist yo'q", "OWASP ASVS L1, HTTPS, throttling, CSP"],
        ["Ishlash", "API javobi lokal < 300 ms; sinxron gunicorn", "p95 < 500 ms, 100 parallel foydalanuvchi"],
        ["Ishonchlilik", "Zaxira va monitoring yo'q", "Kunlik zaxira, Sentry, uptime 99.5%"],
        ["Fayllar", "Bitta fayl 10 MB, so'rov 12 MB; rasm sifati 80%", "S3/obyekt saqlash, rasmni siqish"],
        ["Til va vaqt", "O'zbek (lotin), Asia/Tashkent", "+ rus, o'zbek (kirill)"],
        ["Moslik", "Android, iOS 13+, Web (Chrome); panellar zamonaviy brauzerlar", "Android 8+, iOS 14+"],
        ["Testlar", "99 avtomatik test (backend + bot); frontend testlari yo'q", "Asosiy oqimlar uchun e2e testlar"],
    ],
    widths=[3.0, 7.5, 6.5],
)

# ================================================================= 14
h1("14. Infratuzilma, sozlamalar va ishga tushirish")
h2("14.1. Lokal ishga tushirish (Windows)")
table(
    ["Qism", "Buyruq"],
    [
        ["Backend (birinchi marta)", "powershell -ExecutionPolicy Bypass -File .\\run.ps1 -Setup"],
        ["Backend (API)", ".\\run.ps1 -ApiOnly -NoMigrate   → http://127.0.0.1:8000"],
        ["Demo ma'lumot", ".\\.venv\\Scripts\\python.exe manage.py seed_demo"],
        ["Superadmin paneli", "cd superadmin_panel; npm run dev   → http://localhost:3000"],
        ["Firma paneli", "cd admin_panel; npm run dev   → http://localhost:3002"],
        ["Mobil (Chrome)", "flutter run -d chrome --dart-define=API_BASE_URL=http://127.0.0.1:8000/api/v1"],
        ["Testlar", ".\\.venv\\Scripts\\python.exe manage.py test"],
    ],
    widths=[4.0, 13.0],
)
para("Eslatma: Windows foydalanuvchi papkasida bo'sh joy bo'lgani uchun Flutter qisqa yo'l (8.3) va PUB_CACHE bilan ishga tushiriladi. "
     "Bu kompyuterda Android SDK va Visual Studio yo'q, shuning uchun mobil ilova hozircha faqat Chrome'da ishga tushadi.")
h2("14.2. Asosiy muhit o'zgaruvchilari (.env)")
table(
    ["Guruh", "O'zgaruvchilar"],
    [
        ["Django", "DJANGO_SECRET_KEY, DJANGO_DEBUG, DJANGO_ALLOWED_HOSTS, DJANGO_SETTINGS_MODULE"],
        ["Baza", "USE_SQLITE, POSTGRES_DB, POSTGRES_USER, POSTGRES_PASSWORD, POSTGRES_HOST, POSTGRES_PORT"],
        ["CORS", "CORS_ALLOWED_ORIGINS"],
        ["OTP", "OTP_DEBUG_RETURN_CODE, OTP_CODE_TTL_SECONDS, OTP_CODE_LENGTH, OTP_MAX_ATTEMPTS, OTP_REQUEST_RATE_LIMIT, OTP_REQUEST_RATE_WINDOW_SECONDS"],
        ["SMS", "SMS_PROVIDER, SMS_API_URL, SMS_API_TOKEN"],
        ["To'lov", "PAYME_MERCHANT_ID, PAYME_CHECKOUT_URL, PAYME_ACCOUNT_FIELD, CLICK_SERVICE_ID, CLICK_MERCHANT_ID, CLICK_CHECKOUT_URL, PAYMENT_RETURN_URL, PAYMENTS_TEST_MODE"],
        ["Redis", "REDIS_HOST, REDIS_PORT, REDIS_DB, REDIS_URL, NOTIFICATION_QUEUE_KEY"],
        ["Bot", "BOT_TOKEN, EMAKON_API_BASE_URL, EMAKON_BOT_SERVICE_KEY, BOT_SECRET_KEY, BOT_DATABASE_PATH, BOT_RUN_MODE, BOT_USE_REDIS, TELEGRAM_PROXY, WEBHOOK_URL, WEBHOOK_SECRET"],
        ["Panellar", "NEXT_PUBLIC_API_BASE_URL, NEXT_PUBLIC_GOOGLE_MAPS_KEY (firma paneli)"],
        ["Mobil", "--dart-define API_BASE_URL, USE_LOCAL_DATA"],
    ],
    widths=[2.5, 14.5],
    font_size=8.5,
)
h2("14.3. Production uchun kerakli arxitektura (maqsad)")
bullets(
    [
        "nginx (HTTPS, statik va media fayllar) → gunicorn (Django API, 3+ worker) → PostgreSQL 16 (kunlik zaxira).",
        "Redis (parol bilan) → Telegram bot (webhook yoki polling) va Celery worker/beat (rejalashtirilgan vazifalar).",
        "Media: S3-mos obyekt saqlash; Sentry; uptime monitoring; CI/CD orqali avtomatik deploy.",
    ]
)

# ================================================================= 15
h1("15. Qabul qilish mezonlari (production versiya uchun)")
numbered(
    [
        "Mijoz real telefon raqamiga SMS kod olib kiradi (debug kod o'chiq).",
        "Mijoz xizmat va firmani tanlaydi, faqat yashil (bo'sh) vaqtni tanlay oladi; band vaqtga server ruxsat bermaydi.",
        "Click yoki Payme orqali to'lov avtomatik tasdiqlanadi va escrow'da «held» bo'ladi.",
        "Firma buyurtmani qabul qiladi, xodim biriktiradi, bosqichlar mijozga real vaqtda (push) ko'rinadi.",
        "Ish yakunlangach pul komissiya ayirilib firmaga o'tadi, ledger'da to'g'ri yozuvlar bor.",
        "Nizo/bekor qilishda pul mijozga qaytariladi, jarima va savdo taqiqi ishlaydi.",
        "Superadmin firmalarni moderatsiya qiladi, hisobot va PDF oladi.",
        "Barcha avtomatik testlar CI'da yashil; manage.py check --deploy ogohlantirishsiz.",
        "Mobil ilova Google Play va App Store'da tasdiqlangan.",
    ]
)

# ================================================================= Ilova A
page_break()
h1("Ilova A. API endpointlar ro'yxati (/api/v1/)")
h2("A.1. Mijoz va ommaviy")
table(
    ["Metod va yo'l", "Kirish", "Vazifasi"],
    [
        ["POST auth/otp/request/", "Ochiq", "SMS kod so'rash"],
        ["POST auth/otp/verify/", "Ochiq", "Kodni tekshirish, JWT olish"],
        ["POST auth/admin/login/", "Ochiq", "Admin/superadmin parol bilan kirish"],
        ["POST auth/token/refresh/", "Ochiq", "Tokenni yangilash"],
        ["GET/PATCH auth/me/", "Kirgan", "Profil"],
        ["POST auth/telegram/link/, unlink/", "Mijoz + bot kaliti", "Telegram bog'lash"],
        ["GET services/, services/{slug}/, services/{slug}/offers/", "Ochiq", "Katalog va firma takliflari"],
        ["GET partners/, partners/{id}/", "Ochiq", "Faol firmalar (?service=)"],
        ["GET partners/{id}/availability/?date=", "Ochiq", "Bo'sh/band soatlar"],
        ["GET/POST partners/{id}/reviews/", "Ochiq / mijoz", "Sharhlar"],
        ["GET banners/", "Ochiq", "Faol bannerlar"],
        ["GET/POST orders/, GET orders/{id}/", "Mijoz", "Buyurtmalar (Idempotency-Key)"],
        ["POST orders/{id}/cancel/, pay/, payment-sent/, test-pay/", "Mijoz", "Bekor qilish va to'lov"],
        ["GET care-contracts/", "Mijoz", "Parvarish shartnomalari (o'qish)"],
        ["GET/POST support/, POST support/{id}/reply/", "Mijoz", "Murojaatlar"],
        ["GET chats/, POST chats/open/, GET chats/{id}/messages/, POST chats/{id}/send/", "Mijoz", "Firma bilan chat"],
        ["GET loyalty/, loyalty/transactions/, POST loyalty/redeem/", "Mijoz", "Bonus ballar"],
        ["GET notifications/, POST {id}/read/, read-all/", "Kirgan", "Bildirishnomalar"],
    ],
    widths=[7.5, 3.0, 6.5],
    font_size=8.5,
)
h2("A.2. Admin (firma admini va superadmin), /api/v1/admin/")
table(
    ["Metod va yo'l", "Kirish", "Vazifasi"],
    [
        ["GET dashboard/, reports/bundle/, map/", "Admin (analitika)", "Ko'rsatkichlar, hisobot, xarita"],
        ["GET/POST orders/, GET orders/{id}/", "Admin (buyurtmalar)", "Buyurtmalar ro'yxati"],
        ["POST orders/{id}/transition/, stage/, schedule/", "Admin", "Holat, bosqich, vaqt va uzaytirish"],
        ["GET orders/availability/?date=&organization=", "Admin", "Bandlik jadvali"],
        ["POST orders/{id}/finance/{amal}/ — ensure, mark_paid, reject_payment, release, refund, dispute, punish_firm, punish_user", "Superadmin", "Escrow amallari"],
        ["CRUD services/, POST services/{id}/approve/, reject/", "Admin / superadmin", "Xizmatlar va moderatsiya"],
        ["GET firms/, firms/{id}/stats/, ledger/; GET/PATCH firms/me/", "Admin", "Firma ma'lumotlari"],
        ["POST firms/{id}/{amal}/ — set_trial, apply_suggested_rate, adjust_debt, block, unblock, end_agreement, send_message, add_fine, ban_sales, lift_sales_ban", "Superadmin", "Firma boshqaruvi"],
        ["CRUD employees/, GET {id}/card/, POST {id}/set-status/", "Admin (xodimlar)", "Xodimlar"],
        ["CRUD admins/", "Superadmin", "Firma adminlari"],
        ["GET/POST customers/, segments/, POST {id}/block/, unblock/", "Admin / superadmin", "Mijozlar"],
        ["CRUD care-contracts/ + submit, approve, reject, pause, resume, complete, cancel, comment, visits/{vid}/", "Admin / superadmin", "Parvarish"],
        ["CRUD support/, POST {id}/reply/; chats/…", "Admin", "Murojaat va chat"],
        ["CRUD banners/, loyalty-rewards/; GET loyalty-transactions/; GET/PATCH loyalty/settings/", "Superadmin", "Banner va ball tizimi"],
        ["GET/POST investors/, POST {id}/end_agreement/", "Superadmin", "Investorlar"],
    ],
    widths=[9.0, 3.3, 4.7],
    font_size=8.5,
)

# ================================================================= Ilova B
h1("Ilova B. Ma'lumotlar modeli (asosiy jadvallar)")
table(
    ["Model", "Asosiy maydonlar"],
    [
        ["User", "phone (login), full_name, role (customer/worker/admin/superadmin), organization, loyalty_points, telegram_id, manzil va koordinatalar"],
        ["OTPChallenge", "phone, code_hash, purpose, attempts, is_used, expires_at"],
        ["Organization", "name, slug, status, commission_rate, subscription_plan/units, rating, debt_amount, unpaid_fines, sales_banned_until, trial_ends_at, work_hours, koordinatalar"],
        ["FirmReview / FirmMessage / FirmFine / FirmModerationLog / Investor", "Sharhlar, xabarlar, jarimalar, moderatsiya jurnali, investorlar"],
        ["Service", "organization, base_service (katalog turi), name, slug, price_from/to, moderation_status, is_active, rasmlar"],
        ["Banner", "status, placement, link_service, starts_at, ends_at, image"],
        ["Order", "customer, organization, service, status, work_stage, quoted_price, platform_share, scheduled_date, scheduled_start, duration_minutes, assigned_worker, manzil, koordinatalar, ETA"],
        ["OrderMedia / OrderStatusHistory / OrderIdempotency", "Rasmlar/videolar, holatlar tarixi, takroriy so'rov himoyasi"],
        ["OrderEscrow / OrderPayment / LedgerEntry", "Escrow holati va summalar, to'lov urinishlari, buxgalteriya yozuvlari"],
        ["EmployeeProfile / AdminProfile", "Xodim (mutaxassislik, holat, reyting), admin (imkoniyat bayroqlari)"],
        ["CareContract / CareVisit / CareContractEvent", "Shartnoma, tashriflar, hodisalar jurnali"],
        ["SupportTicket / SupportMessage / ChatRoom / ChatMessage", "Murojaatlar va chat"],
        ["LoyaltySettings / LoyaltyReward / PointTransaction", "Ball qoidalari, mukofotlar, tranzaksiyalar"],
        ["UserNotification / NotificationDelivery", "Ilova ichidagi xabarlar, Telegram yetkazish jurnali"],
    ],
    widths=[5.5, 11.5],
    font_size=8.5,
)

# ================================================================= Ilova C
h1("Ilova C. Demo kirish ma'lumotlari (faqat lokal muhit)")
table(
    ["Qayerda", "Login", "Parol / kod"],
    [
        ["Superadmin paneli — http://localhost:3000", "+998900000001", "admin12345"],
        ["Firma paneli — http://localhost:3002 (Yashil Bog' Servis)", "+998900000002", "firma12345"],
        ["Firma paneli — http://localhost:3002 (Gulzor Landshaft)", "+998900000003", "firma12345"],
        ["Mobil ilova (mijoz)", "+998901000101", "SMS kod ekranda chiqadi (debug rejim)"],
    ],
    widths=[8.0, 4.0, 5.0],
)
para("Diqqat: bu parollar faqat demo uchun. Production'da seed_demo ishlatilmaydi va barcha parollar almashtiriladi.", color=RED)

add_page_number_footer()
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("saved", OUT)
